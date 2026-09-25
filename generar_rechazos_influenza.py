import pandas as pd
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def clean_run(run_str):
    if pd.isna(run_str) or str(run_str).strip() == '':
        return ""
    s = str(run_str).strip().upper()
    if '-' in s:
        s = s.split('-')[0]
    s = s.replace('.', '').replace('-', '')
    return s

def generar_excel(df_ocurrencia, df_residencia, output_path, max_date_str):
    wb = openpyxl.Workbook()
    
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
    
    headers_rechazos = [
        "RUN", "Nombres", "Apellido Paterno", "Apellido Materno",
        "Fecha Rechazo", "Comuna Ocurrencia", "Comuna Residencia", "Establecimiento", "Grupo Objetivo", "Causa Rechazo",
        "RUN Vacunador", "Nombre Vacunador", "RUN Registrador", "Nombre Registrador"
    ]

    # --- Hoja 0: Resumen Ocurrencia ---
    ws_res = wb.active
    ws_res.title = "RESUMEN OCURRENCIA"
    ws_res.sheet_view.showGridLines = False

    # Estilos del resumen
    fill_titlebox = PatternFill(start_color="FCF6D6", end_color="FCF6D6", fill_type="solid")
    font_titlebox = Font(name="Times New Roman", size=11, bold=True)
    
    fill_header = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    font_header = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    
    fill_comuna = PatternFill(start_color="8EAADB", end_color="8EAADB", fill_type="solid")
    font_comuna = Font(name="Tahoma", size=10, bold=True, color="FFFFFF")
    
    font_estab = Font(name="Tahoma", size=10)
    font_total = Font(name="Tahoma", size=10, bold=True)
    border_thick_bottom = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thick"))

    ws_res.column_dimensions['A'].width = 60
    ws_res.column_dimensions['B'].width = 15

    ws_res.merge_cells('A2:B4')
    cell = ws_res['A2']
    cell.value = "Auditoría Nominal de Rechazos Influenza 2026.\nResumen por Ocurrencia."
    cell.font = font_titlebox
    cell.fill = fill_titlebox
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
    ws_res['B2'].border = thin_border
    ws_res['A3'].border = thin_border
    ws_res['B3'].border = thin_border
    ws_res['A4'].border = thin_border
    ws_res['B4'].border = thin_border

    ws_res.merge_cells('A6:B6')
    cell = ws_res['A6']
    cell.value = "SS. Osorno"
    cell.font = Font(name="Times New Roman", size=11, bold=True)
    cell.alignment = Alignment(horizontal="center")

    ws_res.merge_cells('A8:B8')
    cell = ws_res['A8']
    cell.value = f"Fuente DEIS-MINSAL con fecha de extracción {max_date_str}"
    cell.font = Font(name="Times New Roman", size=10, italic=True)
    cell.alignment = Alignment(horizontal="center")

    cell = ws_res.cell(row=11, column=1, value="COMUNA / ESTABLECIMIENTO:")
    cell.fill = fill_header
    cell.font = font_header
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="left", vertical="center")
    
    cell = ws_res.cell(row=11, column=2, value="N°")
    cell.fill = fill_header
    cell.font = font_header
    cell.border = thin_border
    cell.alignment = Alignment(horizontal="center", vertical="center")

    # Datos resumen
    summary_data = []
    total_general = len(df_ocurrencia)
    
    df_ocurrencia_sum = df_ocurrencia.copy()
    df_ocurrencia_sum['COMUNA_OCURR'] = df_ocurrencia_sum['COMUNA_OCURR'].fillna('SIN COMUNA').astype(str)
    df_ocurrencia_sum['ESTABLECIMIENTO'] = df_ocurrencia_sum['ESTABLECIMIENTO'].fillna('SIN ESTABLECIMIENTO').astype(str)
    
    comunas = sorted(df_ocurrencia_sum['COMUNA_OCURR'].unique())
    for comuna in comunas:
        df_comuna = df_ocurrencia_sum[df_ocurrencia_sum['COMUNA_OCURR'] == comuna]
        total_comuna = len(df_comuna)
        summary_data.append({'COMUNA_ESTABLECIMIENTO': f"- {comuna}", 'TOTAL': total_comuna, 'IS_COMUNA': True})
        
        estab_counts = df_comuna['ESTABLECIMIENTO'].value_counts()
        for estab, count in estab_counts.items():
            summary_data.append({'COMUNA_ESTABLECIMIENTO': f"    {estab}", 'TOTAL': count, 'IS_COMUNA': False})

    row_idx = 12
    if summary_data:
        for item in summary_data:
            cell_lbl = ws_res.cell(row=row_idx, column=1, value=item['COMUNA_ESTABLECIMIENTO'])
            cell_num = ws_res.cell(row=row_idx, column=2, value=item['TOTAL'])
            
            cell_lbl.border = thin_border
            cell_num.border = thin_border
            
            if item['IS_COMUNA']:
                cell_lbl.fill = fill_comuna
                cell_lbl.font = font_comuna
                cell_lbl.alignment = Alignment(horizontal="left")
                
                cell_num.fill = fill_comuna
                cell_num.font = font_comuna
                cell_num.alignment = Alignment(horizontal="center")
            else:
                cell_lbl.font = font_estab
                cell_lbl.alignment = Alignment(horizontal="left")
                
                cell_num.font = font_estab
                cell_num.alignment = Alignment(horizontal="center")
            row_idx += 1
    else:
        cell_lbl = ws_res.cell(row=row_idx, column=1, value="Sin información de comuna/establecimiento")
        cell_lbl.font = font_estab
        cell_lbl.border = thin_border
        cell_num = ws_res.cell(row=row_idx, column=2, value=total_general)
        cell_num.font = font_estab
        cell_num.border = thin_border
        cell_num.alignment = Alignment(horizontal="center")
        row_idx += 1

    cell_lbl = ws_res.cell(row=row_idx, column=1, value="Total general")
    cell_lbl.font = font_total
    cell_lbl.border = border_thick_bottom
    cell_num = ws_res.cell(row=row_idx, column=2, value=total_general)
    cell_num.font = font_total
    cell_num.border = border_thick_bottom
    cell_num.alignment = Alignment(horizontal="center")
    
    # --- Hoja 1: Ocurrencia ---
    ws1 = wb.create_sheet(title="Auditoria por Ocurrencia")
    
    for col_idx, header in enumerate(headers_rechazos, 1):
        cell = ws1.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    df_ocurrencia_sorted = df_ocurrencia.sort_values(by=['COMUNA_OCURR', 'ESTABLECIMIENTO', 'NOMBRES'])
    
    for _, row in df_ocurrencia_sorted.iterrows():
        row_idx_data = ws1.max_row + 1
        ws1.cell(row=row_idx_data, column=1, value=row.RUN).border = thin_border
        ws1.cell(row=row_idx_data, column=2, value=row.NOMBRES).border = thin_border
        ws1.cell(row=row_idx_data, column=3, value=row.APELLIDO_PATERNO).border = thin_border
        ws1.cell(row=row_idx_data, column=4, value=row.APELLIDO_MATERNO).border = thin_border
        ws1.cell(row=row_idx_data, column=5, value=row.FECHA_INMUNIZACION).border = thin_border
        ws1.cell(row=row_idx_data, column=6, value=row.COMUNA_OCURR).border = thin_border
        ws1.cell(row=row_idx_data, column=7, value=row.COMUNA_RESIDENCIA).border = thin_border
        ws1.cell(row=row_idx_data, column=8, value=row.ESTABLECIMIENTO).border = thin_border
        ws1.cell(row=row_idx_data, column=9, value=row.CRITERIO_ELEGIBILIDAD).border = thin_border
        ws1.cell(row=row_idx_data, column=10, value=row.CAUSA_RECHAZO).border = thin_border
        ws1.cell(row=row_idx_data, column=11, value=row.FUNCIONARIO_INMUNIZADOR_RUN).border = thin_border
        ws1.cell(row=row_idx_data, column=12, value=row.FUNCIONARIO_INMUNIZADOR).border = thin_border
        ws1.cell(row=row_idx_data, column=13, value=row.FUNCIONARIO_REGISTRADOR_RUN).border = thin_border
        ws1.cell(row=row_idx_data, column=14, value=row.FUNCIONARIO_REGISTRADOR).border = thin_border
        
    for col_idx in range(1, len(headers_rechazos) + 1):
        ws1.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 25
    
    # --- Hoja 2: Residencia ---
    ws2 = wb.create_sheet(title="Auditoria por Residencia")
    
    for col_idx, header in enumerate(headers_rechazos, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin_border
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    df_residencia_sorted = df_residencia.sort_values(by=['COMUNA_RESIDENCIA', 'ESTABLECIMIENTO', 'NOMBRES'])
    
    for _, row in df_residencia_sorted.iterrows():
        row_idx_data = ws2.max_row + 1
        ws2.cell(row=row_idx_data, column=1, value=row.RUN).border = thin_border
        ws2.cell(row=row_idx_data, column=2, value=row.NOMBRES).border = thin_border
        ws2.cell(row=row_idx_data, column=3, value=row.APELLIDO_PATERNO).border = thin_border
        ws2.cell(row=row_idx_data, column=4, value=row.APELLIDO_MATERNO).border = thin_border
        ws2.cell(row=row_idx_data, column=5, value=row.FECHA_INMUNIZACION).border = thin_border
        ws2.cell(row=row_idx_data, column=6, value=row.COMUNA_OCURR).border = thin_border
        ws2.cell(row=row_idx_data, column=7, value=row.COMUNA_RESIDENCIA).border = thin_border
        ws2.cell(row=row_idx_data, column=8, value=row.ESTABLECIMIENTO).border = thin_border
        ws2.cell(row=row_idx_data, column=9, value=row.CRITERIO_ELEGIBILIDAD).border = thin_border
        ws2.cell(row=row_idx_data, column=10, value=row.CAUSA_RECHAZO).border = thin_border
        ws2.cell(row=row_idx_data, column=11, value=row.FUNCIONARIO_INMUNIZADOR_RUN).border = thin_border
        ws2.cell(row=row_idx_data, column=12, value=row.FUNCIONARIO_INMUNIZADOR).border = thin_border
        ws2.cell(row=row_idx_data, column=13, value=row.FUNCIONARIO_REGISTRADOR_RUN).border = thin_border
        ws2.cell(row=row_idx_data, column=14, value=row.FUNCIONARIO_REGISTRADOR).border = thin_border
        
    for col_idx in range(1, len(headers_rechazos) + 1):
        ws2.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = 25
        
    wb.save(output_path)

def process_file_for_excel(csv_path, comunas_osorno, filter_col):
    print(f"\nLeyendo base de vacunacion: {csv_path}")
    cols = ['VACUNA_ADMINISTRADA', 'REGISTRO_ELIMINADO', 'CRITERIO_ELEGIBILIDAD', 'DOSIS', 'CAUSA_RECHAZO', 'RUN', 'FECHA_INMUNIZACION', 'COMUNA_RESIDENCIA', 'COMUNA_OCURR', 'ESTABLECIMIENTO', 'NOMBRES', 'APELLIDO_PATERNO', 'APELLIDO_MATERNO', 'FUNCIONARIO_INMUNIZADOR_RUN', 'FUNCIONARIO_INMUNIZADOR', 'FUNCIONARIO_REGISTRADOR_RUN', 'FUNCIONARIO_REGISTRADOR']
    try:
        df = pd.read_csv(csv_path, sep="|", encoding="utf-8", usecols=cols, dtype=str)
    except UnicodeDecodeError:
        df = pd.read_csv(csv_path, sep="|", encoding="latin-1", usecols=cols, dtype=str)
        
    print(f"Total registros originales CSV: {len(df)}")
    
    # Cleaning columns for filtering
    df['VACUNA_ADMINISTRADA'] = df['VACUNA_ADMINISTRADA'].fillna("").astype(str).str.strip().str.upper()
    df['REGISTRO_ELIMINADO'] = df['REGISTRO_ELIMINADO'].fillna("NO").astype(str).str.strip().str.upper()
    df['CRITERIO_ELEGIBILIDAD'] = df['CRITERIO_ELEGIBILIDAD'].fillna("").astype(str).str.strip()
    df['DOSIS'] = df['DOSIS'].fillna("").astype(str).str.strip().str.upper()
    df['COMUNA_OCURR'] = df['COMUNA_OCURR'].fillna('DESCONOCIDA').astype(str).str.strip().str.upper()
    df['COMUNA_RESIDENCIA'] = df['COMUNA_RESIDENCIA'].fillna('DESCONOCIDA').astype(str).str.strip().str.upper()
    
    # Filter by Osorno communes
    df = df[df[filter_col].isin(comunas_osorno)]
    print(f"Total registros en Osorno ({filter_col}): {len(df)}")

    # 1. Identify pure vaccinated people (to exclude from rejections)
    vacunados_mask = (
        (df['VACUNA_ADMINISTRADA'] == 'SI') & 
        (df['REGISTRO_ELIMINADO'] == 'NO') & 
        (df['CRITERIO_ELEGIBILIDAD'].str.upper() != 'EPRO') & 
        (df['DOSIS'] != 'EPRO')
    )
    df_vacunados = df[vacunados_mask]
    df_vacunados_run = df_vacunados['RUN'].apply(clean_run)
    vacunados_set = set(df_vacunados_run.dropna().unique())
    print(f"Total personas vacunadas: {len(vacunados_set)}")
    
    # 2. Identify rejections
    rechazos_mask = (
        (df['VACUNA_ADMINISTRADA'] != 'SI') & 
        (df['REGISTRO_ELIMINADO'] == 'NO') & 
        (df['CRITERIO_ELEGIBILIDAD'].str.upper() != 'EPRO') & 
        (df['DOSIS'] != 'EPRO') & 
        (df['CAUSA_RECHAZO'].notna()) & 
        (df['CAUSA_RECHAZO'].str.strip() != '')
    )
    df_rechazos = df[rechazos_mask].copy()
    print(f"Total registros de rechazo iniciales: {len(df_rechazos)}")
    
    # Clean RUN for deduplication and matching
    df_rechazos['RUN_CLEAN'] = df_rechazos['RUN'].apply(clean_run)
    
    # Exclude those who were eventually vaccinated
    df_rechazos = df_rechazos[~df_rechazos['RUN_CLEAN'].isin(vacunados_set)]
    
    # Deduplicate rejections (if a person has multiple, keep the first)
    df_rechazos = df_rechazos.drop_duplicates(subset=['RUN_CLEAN'], keep='first')
    
    print(f"Total rechazos reales (sin vacuna): {len(df_rechazos)}")
        
    # Drop RUN_CLEAN for final output
    df_rechazos = df_rechazos.drop(columns=['RUN_CLEAN', 'VACUNA_ADMINISTRADA', 'REGISTRO_ELIMINADO', 'DOSIS'])
    
    # Reorder columns
    cols_order = [
        'RUN', 'NOMBRES', 'APELLIDO_PATERNO', 'APELLIDO_MATERNO', 'FECHA_INMUNIZACION', 
        'COMUNA_OCURR', 'COMUNA_RESIDENCIA', 'ESTABLECIMIENTO', 'CRITERIO_ELEGIBILIDAD', 'CAUSA_RECHAZO',
        'FUNCIONARIO_INMUNIZADOR_RUN', 'FUNCIONARIO_INMUNIZADOR', 
        'FUNCIONARIO_REGISTRADOR_RUN', 'FUNCIONARIO_REGISTRADOR'
    ]
    
    # Ensure columns exist in dataframe to prevent KeyError
    existing_cols = [c for c in cols_order if c in df_rechazos.columns]
    df_rechazos = df_rechazos[existing_cols]
    
    return df_rechazos

def main():
    comunas_osorno = [
        'OSORNO', 'PUERTO OCTAY', 'PURRANQUE', 'PUYEHUE', 
        'RÍO NEGRO', 'RIO NEGRO', 'SAN JUAN DE LA COSTA', 'SAN PABLO'
    ]

    csv_path_residencia = r"C:\Antigravity IDE\WEB DEIS\BASE DATOS MINSAL\2026\Influenza_Residencia_2026.csv"
    csv_path_ocurrencia = r"C:\Antigravity IDE\WEB DEIS\BASE DATOS MINSAL\2026\Influenza_Ocurrencia_2026.csv"
    
    excel_dir = r"C:\Antigravity IDE\WEB DEIS\Influenza_Web"
    os.makedirs(excel_dir, exist_ok=True)
    output_path = os.path.join(excel_dir, "Rechazos_Influenza_General_2026.xlsx")
    
    password = "DEIS2026"
    try:
        with open(r"C:\Antigravity_Secrets\rni.env", "r") as f:
            for line in f:
                if "=" in line:
                    k, v = line.strip().split("=", 1)
                    if k == "EXCEL_PASSWORD":
                        password = v.strip('"')
    except Exception:
        pass
    
    print("PROCESANDO BASE RESIDENCIA...")
    df_residencia = process_file_for_excel(csv_path_residencia, comunas_osorno, 'COMUNA_RESIDENCIA')

    print("\nPROCESANDO BASE OCURRENCIA...")
    df_ocurrencia = process_file_for_excel(csv_path_ocurrencia, comunas_osorno, 'COMUNA_OCURR')
    
    try:
        from datetime import datetime
        if 'FECHA_INMUNIZACION' in df_ocurrencia.columns:
            fechas_dt = pd.to_datetime(df_ocurrencia['FECHA_INMUNIZACION'], format="%Y-%m-%d", errors="coerce")
            max_date_str = fechas_dt.max().strftime("%d-%m-%Y")
        else:
            max_date_str = datetime.fromtimestamp(os.path.getmtime(csv_path_ocurrencia)).strftime("%d-%m-%Y")
    except:
        from datetime import datetime
        max_date_str = datetime.now().strftime("%d-%m-%Y")

    print("\nGuardando nomina de rechazos GENERAL en: " + output_path)
    generar_excel(df_ocurrencia, df_residencia, output_path, max_date_str)    
    
    # Generate by comuna
    archivos_generados = [output_path]
    base_dir = os.path.dirname(output_path)
    
    unique_comunas = set(df_ocurrencia['COMUNA_OCURR'].dropna().unique().tolist() + df_residencia['COMUNA_RESIDENCIA'].dropna().unique().tolist())
    for comuna in unique_comunas:
        if comuna not in comunas_osorno: continue
        
        comuna_filename = f"Rechazos_Influenza_Ocurrencia_2026_{comuna}.xlsx"
        comuna_path = os.path.join(base_dir, comuna_filename)
        
        df_ocur_comuna = df_ocurrencia[df_ocurrencia['COMUNA_OCURR'] == comuna]
        df_resi_comuna = df_residencia[df_residencia['COMUNA_RESIDENCIA'] == comuna]
        
        if df_ocur_comuna.empty and df_resi_comuna.empty: continue
        
        print(f"Generando archivo para comuna: {comuna}")
        generar_excel(df_ocur_comuna, df_resi_comuna, comuna_path, max_date_str)
        archivos_generados.append(comuna_path)
        
    try:
        import win32com.client
        print(f"\nCifrando {len(archivos_generados)} archivos con contrasena '{password}'...")
        excel = win32com.client.Dispatch("Excel.Application")
        excel.DisplayAlerts = False
        
        for p in archivos_generados:
            abs_path = os.path.abspath(p)
            try:
                wb = excel.Workbooks.Open(abs_path)
                wb.Password = password
                wb.SaveAs(abs_path, Password=password)
                wb.Close()
                print(f"Archivo cifrado exitosamente: {os.path.basename(p)}")
            except Exception as ex:
                print(f"Error cifrando {p}: {ex}")
                
        excel.Quit()
        print("Cifrado completado!")
    except Exception as e:
        print(f"ADVERTENCIA: No se pudo cifrar el archivo con win32com. Error: {e}")

    # --- QA Territorial ---
    print("\n" + "="*50)
    print(" QA TERRITORIAL DEL EXCEL DE RECHAZOS")
    print("="*50)
    
    print("\n--- HOJA: Auditoria por Ocurrencia ---")
    print(f"Total RUT unicos: {df_ocurrencia['RUN'].nunique()}")
    print("Comunas unicas de ocurrencia:")
    print(sorted(df_ocurrencia['COMUNA_OCURR'].dropna().unique().tolist()))
    print(f"\nComunas unicas de residencia ({df_ocurrencia['COMUNA_RESIDENCIA'].nunique()} comunas):")
    print(sorted(df_ocurrencia['COMUNA_RESIDENCIA'].dropna().unique().tolist()))
    
    print("\n--- HOJA: Auditoria por Residencia ---")
    print(f"Total RUT unicos: {df_residencia['RUN'].nunique()}")
    print("Comunas unicas de residencia:")
    print(sorted(df_residencia['COMUNA_RESIDENCIA'].dropna().unique().tolist()))
    print(f"\nComunas unicas de ocurrencia ({df_residencia['COMUNA_OCURR'].nunique()} comunas):")
    print(sorted(df_residencia['COMUNA_OCURR'].dropna().unique().tolist()))
    print("="*50 + "\n")

if __name__ == '__main__':
    main()
